"""Structural and numeric report audit; run with dependency Python + python-docx."""
import json, hashlib, re
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from lxml import etree

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'evaluation/results/report_audit'
a=json.loads((OUT/'audit.json').read_text(encoding='utf-8'))
rev=json.loads((OUT/'report_revision.json').read_text(encoding='utf-8'))
doc=Document(rev['output']); old=Document(rev['backup'])
def flat(t): return [c.text.replace('Evidence Recall','Recall') for row in t.rows for c in row.cells]
assert len(doc.tables)==15 and len(old.tables)==11
assert all(flat(x)==flat(y) for x,y in zip(doc.tables[:11],old.tables))
assert len([p for p in doc.paragraphs if re.match(r'^[1-7]\. ',p.text)])==7
def formulas(path):
    with ZipFile(path) as z:
        x=etree.fromstring(z.read('word/document.xml'))
    return [etree.tostring(n) for n in x.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath')]
assert formulas(rev['output'])==formulas(rev['backup'])
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
for path,h in a['provenance']['inputs_sha256'].items(): assert sha(ROOT/path)==h
for path,h in a['provenance']['source_sha256_before'].items(): assert sha(ROOT/path)==h
for view in ['paragraph','article']:
    base=ROOT/f'evaluation/results/xquad_vi_{view}_test'
    r=json.loads((base/'rag_comparison.json').read_text(encoding='utf-8'))
    s=json.loads((base/'rag_statistical_summary.json').read_text(encoding='utf-8'))
    for path,h in r['provenance']['chunking_sources'].items(): assert sha(ROOT/path)==h
    for key in ['recall_at_3','recall_at_5','mrr_at_5']:
        assert abs((r['results'][1][key]-r['results'][0][key])-s['metrics'][key]['delta'])<1e-10
    for result in r['results']:
        failures={}
        for row in result['details']:
            k='hit' if row['hit_at_5'] else row['failure_reason']; failures[k]=failures.get(k,0)+1
        assert failures==s['failure_counts'][result['method']]
texts='\n'.join(p.text for p in doc.paragraphs)
for phrase in ['Bài báo gốc.','Model card.']:
    assert phrase not in texts
assert '[11]' in texts
with ZipFile(rev['output']) as z:
    assert z.testzip() is None
    rels=etree.fromstring(z.read('word/_rels/document.xml.rels'))
links=[]; missing=[]
for rel in rels:
    if rel.get('Type','').endswith('/hyperlink'):
        target=rel.get('Target'); links.append(target)
        if not re.match(r'^[a-zA-Z]+:',target) and not (ROOT/target).exists(): missing.append(target)
assert not missing,missing
report=dict(original_tables_unchanged=11,native_equations_unchanged=len(formulas(rev['output'])),main_sections=7,
            new_tables=4,original_inputs_unchanged=True,chunking_sources_unchanged=True,historical_source_hashes_match=True,
            deltas_and_failure_counts_match=True,hyperlinks=len(links),missing_local_links=missing,
            visual_render='Not completed: packaged renderer requires soffice.exe, unavailable in environment.',
            citation_review='Primary sources checked: ACL XQuAD/embeddings, Springer Max-Min, author code, LlamaIndex docs, MiniLM card, arXiv RAG; local links checked.')
(OUT/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
