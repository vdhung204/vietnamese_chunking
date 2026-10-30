"""Revise the existing Word report, preserving its structure and original tables.
Run with the Codex dependency Python (python-docx); diagnostics must exist first.
"""
from pathlib import Path
import json, re, shutil, hashlib
from copy import deepcopy
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'evaluation/results/report_audit'
a=json.loads((OUT/'audit.json').read_text(encoding='utf-8'))
path=ROOT/'BAO_CAO_CAP_NHAT.docx'
backup=Path.home()/'AppData/Local/Temp/utc-rag-word/BAO_CAO_before_review.docx'
backup.parent.mkdir(parents=True,exist_ok=True)
shutil.copy2(path,backup)
d=Document(path)
md=(ROOT/'BAO_CAO.md').read_text(encoding='utf-8')
def f(x,n=2): return f'{x:.{n}f}'.replace('.',',')
def pct(x): return f(100*x)+'%'
def label(view): return 'Context gốc' if view=='paragraph' else 'Ghép dài'
def method(name): return name.replace('Max-Min','Max–Min')

replacements={
'Đóng góp được kiểm chứng ở đây': 'Đề tài xây dựng và đánh giá một chiến lược chunking tiếng Việt kết hợp ràng buộc độ dài khi chọn điểm cắt, relative semantic drop, dependency-aware penalty và dynamic overlap trong cùng pipeline retrieval. Đây là một phương án tích hợp và đánh giá có đối chứng; kết quả không xác lập tính mới tuyệt đối hoặc ưu thế tổng quát so với các phương pháp khác.',
'Luật liên kết câu chỉ tham gia nhánh trên 240.': 'Trong nhánh phân đoạn đối với phần văn bản dài hơn MAX = 240, dependency risk được đưa vào hàm scoring để giảm khả năng chọn boundary phá quan hệ liên câu. Đây là dependency-aware penalty, không phải hard constraint và không bảo đảm tuyệt đối tránh cắt. Nhánh 120–240 không sử dụng penalty này. Overlap được ghép sau khi chọn ranh giới; khi không tìm được ứng viên, phần còn lại được giữ nguyên, nên 240 không phải giới hạn cứng của chunk đầu ra.',
'Trong benchmark một evidence mỗi câu hỏi này': 'Báo cáo dùng thuật ngữ Evidence Recall@3 và Evidence Recall@5 cho các chỉ số recall_at_3 và recall_at_5 trong mã nguồn. Với một gold evidence chính cho mỗi câu hỏi, Evidence Recall@k tương đương tỷ lệ câu hỏi có toàn bộ bằng chứng chuẩn xuất hiện trong ít nhất một chunk đúng nguồn của top-k. Đây là hit rate theo câu hỏi; không phải tỷ lệ mọi tài liệu liên quan được tìm thấy. Ký hiệu Recall trong các công thức mang cùng ý nghĩa này.',
'Để bổ sung một so sánh thực dụng chặt chẽ hơn': 'Phân tích hậu nghiệm tại mục 5.7 áp cùng trần 512/1024 tokenizer token lên top-5 đã lưu, giữ nguyên thứ hạng và chunk. Phân tích này không thay thế thí nghiệm fixed-budget trên toàn bộ danh sách xếp hạng: còn giới hạn tối đa năm ứng viên, có thể để dư ngân sách và chưa có quy tắc phân bổ context được đăng ký trước. Các cấu hình bổ sung không được dùng để chọn lại tham số hoặc sửa kết quả test gốc.',
'Số liệu dưới đây lấy từ các CSV': 'Các bảng chính dưới đây giữ nguyên số liệu từ CSV/JSON của lần chạy ngày 30/10/2026 [9, 10]. Nhóm đã tính lại chỉ số từ các kết quả từng câu hỏi để kiểm tra tính nhất quán. Các phân tích bổ sung ở mục 5.7–5.9 được tính riêng từ artifact đã lưu [11], không chạy lại chunking hoặc thay dữ liệu test.',
'Trên context gốc, Ours đạt kết quả truy xuất tốt hơn': 'Trong cấu hình top-5 chunk hiện tại, Ours có Evidence Recall@5 cao hơn trên context gốc, nhưng đồng thời sử dụng lượng context lớn hơn đáng kể. Có 81 câu chỉ Ours hit và 34 câu chỉ Max–Min hit, chênh lệch ròng 47 câu. MRR@5 cũng cao hơn. Nhiều context vốn ngắn khiến Ours có thể giữ phần lớn hoặc toàn bộ context; chưa thể quy mức tăng Recall hoàn toàn cho chất lượng xác định boundary. Kết quả dưới trần token ở mục 5.7 cho thấy cần xét cả ngân sách context khi so sánh.',
'Một số cơ chế có thể giải thích kết quả': 'Giả thuyết giải thích: chunk lớn có thể làm vector tổng hợp ít tập trung vào chi tiết câu hỏi; overlap có thể làm thay đổi nội dung vector hoặc tạo lặp trong context. Phân tích tokenizer tại mục 5.8 đo được nhiều chunk Ours vượt giới hạn 128 token của cấu hình SentenceTransformer đang lưu. Truncation là một cơ chế có thể góp phần giải thích kết quả retrieval trên chunk dài, nhưng chưa được chứng minh là nguyên nhân của chênh lệch giữa hai phương pháp. Chưa đo vị trí evidence trong vùng encode và chưa có đối chứng loại trừ truncation.',
'Các thí nghiệm tiếp theo cần gồm:': 'Hướng thực nghiệm tiếp theo: fixed context-token budget trên toàn bộ ranking; phân tích offset evidence so với vùng tokenizer được encode; dependency challenge set có văn bản trên 240 đơn vị; ablation bỏ overlap, bỏ penalty và chỉ dùng độ dài; cuối cùng là đánh giá câu trả lời bằng cùng LLM, prompt và rubric/EM/F1. Các thử nghiệm cần được định nghĩa trước, dùng dev để chọn tham số và xác nhận trên dữ liệu độc lập mới sau khi đã xem test hiện tại.',
'Kết luận phù hợp là hiệu quả của chiến lược': 'Ours chưa chứng minh ưu thế tổng quát so với Max–Min. Lợi thế về bảo toàn evidence và retrieval trên context gốc trong cấu hình top-5 đi kèm lượng context lớn hơn, trong khi lợi thế retrieval đảo chiều trên văn bản ghép dài. Phân tích trần token và tokenizer bổ sung cho thấy kết luận phụ thuộc chính sách context; chưa xác lập quan hệ nhân quả giữa kích thước chunk, truncation và thứ hạng. Hướng tiếp theo là fixed-budget trên toàn ranking, evidence truncation theo offset, dependency diagnostic ở nhánh trên MAX, ablation và đánh giá sinh câu trả lời đầu cuối.',
}
def setpara(p,text):
    p.clear(); p.add_run(text)
for p in d.paragraphs:
    for prefix,text in replacements.items():
        if p.text.startswith(prefix): setpara(p,text); break
    if p.text.startswith('Kiss, Nagy và Szilágyi công bố'):
        setpara(p,p.text.replace('Đây là bằng chứng trong thiết lập của tác giả, chưa xác nhận hiệu quả trên tiếng Việt. Bài báo gốc.','Đây là kết quả trong thiết lập của tác giả [2], chưa trực tiếp xác nhận hiệu quả trên tiếng Việt.'))
    elif p.text.startswith('LlamaIndex cung cấp'):
        setpara(p,p.text.replace('Tài liệu chính thức.','Nguồn triển khai được dẫn tại [4].'))
    elif p.text.startswith('Mô hình sentence-transformers/'):
        setpara(p,p.text.replace('Model card.','Thông số mô hình được dẫn tại [6].'))
for prefix,text in replacements.items():
    pattern=r'(?m)^'+re.escape(prefix)+r'.*$'
    # Original Markdown has bold lead-ins on a few paragraphs.
    lines=md.splitlines()
    for i,line in enumerate(lines):
        if line.replace('**','').startswith(prefix): lines[i]=text
    md='\n'.join(lines)+'\n'
md=md.replace('Đây là bằng chứng trong thiết lập của tác giả, chưa xác nhận hiệu quả trên tiếng Việt. [Bài báo gốc](https://doi.org/10.1007/s10791-025-09638-7).','Đây là kết quả trong thiết lập của tác giả [2], chưa trực tiếp xác nhận hiệu quả trên tiếng Việt.')
md=md.replace('[Tài liệu chính thức](https://developers.llamaindex.ai/python/examples/node_parsers/semantic_chunking/).','Nguồn triển khai được dẫn tại [4].',1)
md=md.replace('[Model card](https://huggingface.co/sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2).','Thông số mô hình được dẫn tại [6].',1)

dev='Vai trò dev và test: các tham số của hai chiến lược được giữ cố định khi chạy XQuAD test, theo README và cấu hình mã nguồn hiện có. Dev gồm 233 câu hỏi được tách riêng để phục vụ phát triển và kiểm tra pipeline; không có artifact ghi nhận một đợt tìm tham số trên XQuAD dev, nên báo cáo không tuyên bố đã tune bằng dev. Test gồm 957 câu hỏi không được dùng để chọn lại threshold/weight trong lần rà soát này. Phân tích bổ sung là hậu nghiệm sau khi đã xem kết quả test, không phải xác nhận độc lập. Hash mã nguồn hiện tại khớp các hash lưu ở cả hai lần chạy XQuAD.'
anchor=next(p for p in d.paragraphs if p.text.startswith('Ở góc nhìn paragraph'))
anchor.insert_paragraph_before(dev,style='Body Text')
md=md.replace('Ở góc nhìn `paragraph`',dev+'\n\nỞ góc nhìn `paragraph`',1)

blocks=[]
def para(text): blocks.append(('p',text))
def heading(text): blocks.append(('h',text))
def table(headers,rows,widths): blocks.append(('t',(headers,rows,widths)))
heading('5.7. Phân tích trần token trên top-5 đã lưu')
para('Thiết kế hậu nghiệm: dùng tokenizer multilingual MiniLM ở snapshot cục bộ e8f8c211226b894fcb81acc59f3b34ba3efd5f42. Với mỗi câu hỏi, duyệt top-5 đã lưu theo thứ hạng gốc, nối nguyên chunk bằng hai ký tự xuống dòng và dừng trước chunk đầu tiên làm vượt trần 512 hoặc 1024 token. Không cắt chunk, không bỏ qua chunk quá dài để lấy chunk sau, không bổ sung ứng viên ngoài top-5 và không dùng evidence để chọn context. Token budget chỉ tính nội dung context, không gồm special token, câu hỏi hoặc prompt; đây là tokenizer của embedding model, chưa phải ngân sách của một LLM cụ thể.')
para('Evidence Recall@budget là tỷ lệ câu hỏi có chunk hợp lệ trong phần context được giữ; MRR@budget dùng nghịch đảo hạng của chunk hợp lệ đầu tiên, bằng 0 nếu không có. Vì chỉ giữ tiền tố của top-5, thứ hạng các chunk được chọn không đổi. Cùng trần không có nghĩa cùng số token thực dùng; giới hạn top-5 có thể ngăn Max–Min sử dụng hết ngân sách. Kết quả là chẩn đoán “top-5 có trần token”, không phải phép so sánh fixed-budget đầy đủ trên toàn ranking. Không tính khoảng tin cậy mới và không chọn mức budget theo kết quả thắng/thua.')
rows=[]
for b in a['budget_results']:
    rows.append([label(b['view']),method(b['method']),str(b['budget']),pct(b['evidence_recall']),f(b['mrr'],4),f(b['mean_context_tokens'])])
table(['Đầu vào','Chiến lược','Trần token','Evidence Recall@budget','MRR@budget','Token thực dùng TB'],rows,[3.1,2.1,1.6,3.6,2.6,3.8])
para('Ở trần 512 token, context gốc cho 819/957 hit với Max–Min và 815/957 với Ours; chênh lệch nhỏ này chưa được kiểm định thống kê. Trên văn bản ghép dài, số hit tương ứng là 800 và 609. Ở trần 1024 token, context gốc cho 832 và 873 hit; văn bản ghép dài cho 816 và 748 hit. MRR trên context gốc vẫn cao hơn ở Ours trong hai mức trần, nên không đồng nhất kết luận Recall với mọi khía cạnh xếp hạng. Đây là bằng chứng về độ nhạy với chính sách context trong tập đã quan sát, không chứng minh cơ chế nhân quả hoặc ưu thế tổng quát [11].')
heading('5.8. Phân tích độ dài tokenizer và nguy cơ truncation')
para('Đo độ dài tất cả chunk đã lưu bằng tokenizer snapshot nêu ở mục 5.7, không truncation và có special tokens. Trung bình, trung vị, cực đại và tỷ lệ vượt 128 được tính theo chunk, không theo câu hỏi. sentence_bert_config.json của snapshot ghi max_seq_length=128; tokenizer có thể có model_max_length khác, nên không dùng giới hạn tokenizer chung để thay cho cấu hình SentenceTransformer. Artifact retrieval cũ không cố định revision tokenizer: các số sau mô tả độ dài dưới snapshot cục bộ được ghi hash trong [11].')
table(['Đầu vào','Chiến lược','TB token','Trung vị','Lớn nhất','Chunk >128','Tỷ lệ'],
      [[label(s['view']),method(s['method']),f(s['mean']),f(s['median'],0),str(s['max']),str(s['over_128'])+'/'+str(s['chunks']),pct(s['fraction_over_128'])] for s in a['tokenizer_stats']],
      [3,2.1,2.2,2.1,2.1,2.9,2.4])
para('Kết quả đo cho thấy tỷ lệ chunk Ours vượt 128 token cao hơn Max–Min ở cả hai cách tổ chức đầu vào. Nếu encode theo cấu hình SentenceTransformer 128 token, các chuỗi vượt giới hạn sẽ bị cắt trước khi biểu diễn. Tuy nhiên, độ dài vượt giới hạn không tự cho biết evidence bị mất: chưa có ánh xạ offset bằng chứng sang từng vùng token, đặc biệt khi có overlap hoặc nhiều lần xuất hiện. Vì vậy, không báo tỷ lệ evidence bị truncation và không kết luận truncation chắc chắn gây ra mức giảm Recall. Cần kiểm tra offset và chạy đối chứng có kiểm soát để xác định tác động [11].')
heading('5.9. Chẩn đoán vi phạm ranh giới dependency')
para('Targeted diagnostic evaluation sử dụng annotation có sẵn trong data/benchmark/annotations.json và chunk của bộ pilot đã lưu tại evaluation/results/rag_comparison.json. Có ba ranh giới dương tính: case_001 tại vị trí 0, case_007 và case_011 tại vị trí 4; vị trí được đánh số từ 0. case_012 là tài liệu đối chứng ngoại lệ với danh sách dependency rỗng, không tăng mẫu số. Không tự sinh nhãn dependency từ XQuAD. Metric là số ranh giới dependency bị chọn làm điểm cắt chia cho tổng ba ranh giới được annotate; nó đo quyết định ranh giới, không đo mọi quan hệ ngữ nghĩa hay khả năng overlap sửa mất ngữ cảnh.')
table(['Chiến lược','Ranh giới vi phạm','Ranh giới đã gán nhãn','Dependency Violation Rate'],
      [[method(x['method']),str(x['violations']),str(x['annotated_dependency_boundaries']),pct(x['rate'])] for x in a['dependency_diagnostic']],
      [3.2,3.6,4.2,5.8])
para('Max–Min cắt tại hai ranh giới được gán nhãn của case_007 và case_011; Ours không cắt tại ba ranh giới dương tính. Tuy nhiên, độ dài toàn văn của case_001, case_007, case_011 và case_012 lần lượt chỉ là 16, 228, 205 và 216 đơn vị khoảng trắng. Do đó không testcase nào kích hoạt nhánh trên 240 có dependency penalty. Kết quả 0/3 của Ours phản ánh hành vi của chiến lược đầy đủ trên các mẫu nhỏ, không chứng minh tác dụng riêng của penalty. Bộ pilot đã được dùng phát triển, không độc lập và quá nhỏ để suy rộng; cần thêm văn bản trên MAX, annotation đáng tin cậy và ablation penalty [11].')
heading('5.10. Tổng hợp đánh đổi và khả năng tái lập')
table(['Tiêu chí trong hai thiết lập đã đo','Max–Min','Ours'],[
['Bảo toàn evidence','98,85% ở cả hai','100% ở cả hai'],
['Độ dài chunk và context top-5','Nhỏ hơn','Lớn hơn'],
['Evidence Recall@5 trên context gốc','86,94%','91,85%'],
['Evidence Recall@5 trên văn bản ghép dài','85,27%','80,56%'],
['Luật liên kết tiếng Việt','Không thêm luật riêng','Penalty chỉ ở nhánh trên MAX'],
['Overlap động','Không thêm','Giữ 0–2 câu'],
], [6.8,4.5,5.5])
para('Bảng tổng hợp mô tả hai cấu hình đã quan sát, không xếp hạng phương pháp cho mọi dữ liệu. Evidence preservation và retrieval ranking là hai mục tiêu khác nhau: Ours có evidence_split = 0 nhưng vẫn có 78 ranking_miss ở context gốc và 186 ở văn bản ghép dài. Giả thuyết vector của chunk nhỏ tập trung hơn cần thí nghiệm nhân quả, không được suy trực tiếp từ bảng.')
para('Tái lập phân tích: chạy .\\venv\\Scripts\\python.exe -X utf8 scripts/audit_report_evidence.py từ thư mục gốc. Script chỉ đọc chunk, ranking và nhãn đã lưu, không gọi chunker hoặc LLM. Kết quả nằm tại evaluation/results/report_audit/: audit.json lưu cấu hình, snapshot/hash tokenizer, phiên bản phần mềm, checksum đầu vào và mã nguồn; các CSV lưu thống kê tokenizer, độ dài từng chunk, kết quả theo budget và từng câu hỏi. Không có lấy mẫu ngẫu nhiên nên seed không áp dụng; seed 2026 và 5.000 lượt chỉ thuộc bootstrap gốc.')
para('Các metric gốc được tính lại từ details và đối chiếu với JSON/CSV: số tài liệu, câu hỏi, chunk, Evidence Recall@3/@5, MRR@5, độ dài chunk/context và bảo toàn evidence. Việc bổ sung phân tích không thay threshold, weight, implementation Max–Min, nhãn test hoặc các bảng kết quả gốc. Các phân tích này chưa cung cấp evidence truncation theo offset, full-ranking fixed-budget, ablation hoặc QA generation/EM/F1.')

anchor=next(p for p in d.paragraphs if p.text=='6. Kết luận')
def add_table(headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers)); t.autofit=False
    for cell,s in zip(t.rows[0].cells,headers): cell.text=s
    for row in rows:
        for cell,s in zip(t.add_row().cells,row): cell.text=s
    for col,width in zip(t.columns,widths): col.width=Cm(width)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if ri==0: row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for cell,width in zip(row.cells,widths):
            cell.width=Cm(width); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side); el.set(qn('w:w'),'80'); el.set(qn('w:type'),'dxa'); mar.append(el)
            cell._tc.get_or_add_tcPr().append(mar)
            if ri==0:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E7EEF5'); cell._tc.get_or_add_tcPr().append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.05
                for r in p.runs:
                    r.font.name='Times New Roman'; r.font.size=Pt(10.5); r.bold=(ri==0)
    anchor._p.addprevious(t._tbl)
    return t
newmd=[]
for kind,value in blocks:
    if kind in ['p','h']:
        p=anchor.insert_paragraph_before(value,style='Heading 2' if kind=='h' else 'Body Text')
        p.paragraph_format.widow_control=True
        newmd.append(('### ' if kind=='h' else '')+value)
    else:
        headers,rows,widths=value; add_table(headers,rows,widths)
        newmd.append('\n'.join(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']+['| '+' | '.join(r)+' |' for r in rows]))
md=md.replace('## 6. Kết luận','\n\n'.join(newmd)+'\n\n## 6. Kết luận',1)

# Change labels without touching metric keys, native equations or hyperlinks.
for p in list(d.paragraphs)+[p for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs]:
    for r in p.runs:
        if 'Recall@' in r.text:
            r.text=re.sub(r'(?<!Evidence )\bRecall@', 'Evidence Recall@',r.text)
md=re.sub(r'(?<!Evidence )\bRecall@','Evidence Recall@',md)
ref='Dự án utc_rag. Phân tích hậu nghiệm phục vụ rà soát báo cáo (30/10/2026). Script: scripts/audit_report_evidence.py. Artifact và hướng dẫn: evaluation/results/report_audit/README.md; dữ liệu kết quả: evaluation/results/report_audit/audit.json và các CSV cùng thư mục.'
tail=next(p for p in d.paragraphs if p.text.startswith('Các mục [9–10]'))
refp=tail.insert_paragraph_before('[11] '+ref,style='Body Text')
md=md.replace('Các mục [9–10]', '11. '+ref+'\n\nCác mục [9–10]',1)
# Retain reference hyperlinks; make terse standalone link labels descriptive.
for p in d.paragraphs:
    if p.text.startswith('Lewis, P.') or p.text.startswith('LlamaIndex. Semantic Chunker.'):
        for el in p._p.iter(qn('w:t')):
            if el.text=='Bài báo': el.text='Toàn văn nghiên cứu của Lewis và cộng sự'
            if el.text=='Tài liệu chính thức': el.text='Hướng dẫn triển khai Semantic Chunker của LlamaIndex'

try: d.save(path)
except PermissionError:
    path=ROOT/'BAO_CAO_RA_SOAT.docx'; d.save(path)
(ROOT/'BAO_CAO.md').write_text(md,encoding='utf-8')
(OUT/'report_revision.json').write_text(json.dumps(dict(output=str(path),backup=str(backup),new_sections=['5.7','5.8','5.9','5.10'],
    original_tables_preserved=11,new_tables=4,total_tables=len(d.tables),source_report_backup_sha256=hashlib.sha256(backup.read_bytes()).hexdigest()),ensure_ascii=False,indent=2),encoding='utf-8')
print('Saved',path,'tables:',len(d.tables))
